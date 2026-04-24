# -*- coding: utf-8 -*-
"""Converte extrato analítico FGTS (TXT) em planilha Excel por trabalhador."""

import re
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from dateutil.relativedelta import relativedelta
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

MESES_MAP = {
    "JANEIRO": "01",
    "FEVEREIRO": "02",
    "MARCO": "03",
    "MARÇO": "03",
    "ABRIL": "04",
    "MAIO": "05",
    "JUNHO": "06",
    "JULHO": "07",
    "AGOSTO": "08",
    "SETEMBRO": "09",
    "OUTUBRO": "10",
    "NOVEMBRO": "11",
    "DEZEMBRO": "12",
}

TABLE_COLUMNS = [
    "Competência",
    "Data do Depósito",
    "Situação FGTS",
    "Valor (R$)",
    "Prazo (20 do mês seguinte)",
]

DEP_RE = re.compile(
    r"(\d{2}/\d{2}/\d{4})\s+DEPOSITO(?:\s+EM\s+ATRASO)?\s+([A-ZÇÃÁÉÍÓÚÊÔÕÂÜ]+)\/(\d{4})\s+([\d\.,-]+)",
    re.UNICODE,
)


def _comp_to_date(comp: str) -> date:
    """Converte 'MM/AAAA' em date(AAAA, MM, 1)."""
    mm, yy = comp.split("/")
    return date(int(yy), int(mm), 1)


def _prazo_fgts(comp: str) -> date:
    """Prazo padrão: dia 20 do mês seguinte à competência."""
    base = _comp_to_date(comp)
    prox = base + relativedelta(months=1)
    return date(prox.year, prox.month, 20)


def _parse_header(block_text: str) -> dict:
    m_nome = re.search(r"NOME DO TRABALHADOR.*\n\s*([A-ZÀ-ÜÇ ]+?)\s+\d{3,}", block_text, re.M)
    m_emp = re.search(r"NOME DO EMPREGADOR.*\n\s*([^\n]+?)\s+\d{14}\s*$", block_text, re.M)
    m_insc = re.search(r"INSCRICAO EMPREGADOR\s*\n\s*[^\n]+?\s+(\d{14})\s*$", block_text, re.M)
    m_adm = re.search(r"DTA\.ADM\.\s+SITUACAO CTA\s*\n[^\n]*?(\d{2}/\d{2}/\d{4})", block_text, re.M)
    m_line = re.search(r"DATA DE OPCAO.*\n\s*([0-9/]{10})\s+([0-9/]{10})\s+([0-9/]{10})", block_text, re.M)

    return {
        "nome_trabalhador": m_nome.group(1).strip() if m_nome else "",
        "empregador": m_emp.group(1).strip() if m_emp else "",
        "inscricao_empregador": m_insc.group(1) if m_insc else "",
        "data_admissao": m_adm.group(1) if m_adm else "",
        "data_afastamento": m_line.group(3) if m_line else "",
    }


def _extract_deposits(block_text: str) -> pd.DataFrame:
    deposits = []
    for m in DEP_RE.finditer(block_text):
        data_dep = datetime.strptime(m.group(1), "%d/%m/%Y").date()
        mes_nome = m.group(2).strip()
        ano = int(m.group(3))
        valor = float(m.group(4).strip().replace(".", "").replace(",", "."))

        mes_num = MESES_MAP.get(mes_nome)
        if not mes_num:
            continue

        deposits.append(
            {
                "competencia": f"{mes_num}/{ano}",
                "data": data_dep,
                "valor": valor,
            }
        )

    return pd.DataFrame(deposits)


def _build_competence_table(header: dict, dep_df: pd.DataFrame) -> pd.DataFrame:
    if dep_df.empty:
        raise ValueError("Nenhum recolhimento identificado neste extrato.")

    agg = dep_df.groupby("competencia").agg(
        valor_total=("valor", "sum"),
        datas=("data", lambda s: sorted(set(s))),
    ).reset_index()

    if header.get("data_admissao"):
        adm_date = datetime.strptime(header["data_admissao"], "%d/%m/%Y").date()
        start_comp_date = date(adm_date.year, adm_date.month, 1)
    else:
        first_comp = min(agg["competencia"], key=_comp_to_date)
        start_comp_date = _comp_to_date(first_comp)

    last_comp = max(agg["competencia"], key=_comp_to_date)
    end_comp_date = _comp_to_date(last_comp)

    competencias = []
    cur = start_comp_date
    while cur <= end_comp_date:
        competencias.append(f"{cur.month:02d}/{cur.year}")
        cur += relativedelta(months=1)

    agg_map = {row["competencia"]: row for _, row in agg.iterrows()}

    rows = []
    for comp in competencias:
        prazo = _prazo_fgts(comp)
        if comp in agg_map:
            datas = agg_map[comp]["datas"]
            valor_total = float(agg_map[comp]["valor_total"])
            em_atraso = any(d > prazo for d in datas)
            situacao = "Em Atraso" if em_atraso else "No Prazo"
            data_dep_str = ", ".join(d.strftime("%d/%m/%Y") for d in datas)
            rows.append([comp, data_dep_str, situacao, valor_total, prazo.strftime("%d/%m/%Y")])
        else:
            rows.append([comp, "—", "Não recolhido", 0.0, prazo.strftime("%d/%m/%Y")])

    return pd.DataFrame(rows, columns=TABLE_COLUMNS)


def _split_workers(full_text: str) -> list[str]:
    marker = "FGTS - EXTRATO ANALITICO DO TRABALHADOR"
    parts = full_text.split(marker)
    if len(parts) <= 1:
        return []
    return [marker + p for p in parts[1:]]


def _safe_sheet_name(name: str, idx: int, used: set[str]) -> str:
    base = (name or f"Trabalhador_{idx}").strip()[:31]
    candidate = base
    suffix = 1
    while candidate in used:
        tail = f"_{suffix}"
        candidate = f"{base[:31-len(tail)]}{tail}"
        suffix += 1
    used.add(candidate)
    return candidate


def extrato_fgts_txt_para_excel(txt_path: str, xlsx_path: str = "Extrato_FGTS_Analitico_Processado.xlsx") -> str:
    content = Path(txt_path).read_text(encoding="utf-8", errors="ignore")

    marker = "FGTS - EXTRATO ANALITICO DO TRABALHADOR"
    if marker not in content:
        raise ValueError("O arquivo enviado não segue o formato do extrato analítico do FGTS.")

    blocks = _split_workers(content)
    if not blocks:
        raise ValueError("O texto não segue o formato de extrato analítico padrão da Caixa.")

    wb = Workbook()
    wb.remove(wb.active)
    used_sheet_names: set[str] = set()

    for idx, block in enumerate(blocks, start=1):
        header = _parse_header(block)
        dep_df = _extract_deposits(block)
        df = pd.DataFrame(columns=TABLE_COLUMNS) if dep_df.empty else _build_competence_table(header, dep_df)

        sheet_name = _safe_sheet_name(header.get("nome_trabalhador", ""), idx, used_sheet_names)
        ws = wb.create_sheet(title=sheet_name)

        bold = Font(bold=True)
        ws["A1"], ws["B1"] = "Nome do Trabalhador:", header.get("nome_trabalhador", "")
        ws["A2"] = "Empregador:"
        insc = header.get("inscricao_empregador", "")
        emp = header.get("empregador", "")
        ws["B2"] = f"{emp} ({insc})" if insc else emp
        ws["A3"], ws["B3"] = "Data de Admissão:", header.get("data_admissao", "")
        ws["A4"], ws["B4"] = "Data de Afastamento:", header.get("data_afastamento", "")
        for c in ("A1", "A2", "A3", "A4"):
            ws[c].font = bold

        start_row = 6
        for row in dataframe_to_rows(df, index=False, header=True):
            ws.append(row)

        header_row = start_row
        fill = PatternFill("solid", fgColor="D9E1F2")
        for col in range(1, df.shape[1] + 1):
            cell = ws.cell(row=header_row, column=col)
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.freeze_panes = ws["A7"]
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(df.shape[1])}{header_row + len(df)}"

        currency_format = "#,##0.00"
        for r in range(header_row + 1, header_row + 1 + len(df)):
            ws.cell(row=r, column=4).number_format = currency_format
            ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=3).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=4).alignment = Alignment(horizontal="right")
            ws.cell(row=r, column=5).alignment = Alignment(horizontal="center")

        for col in range(1, df.shape[1] + 1):
            max_len = 0
            for r in range(1, header_row + 1 + len(df)):
                val = ws.cell(row=r, column=col).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 12), 45)

        ws.column_dimensions["A"].width = 14

    wb.save(xlsx_path)
    return xlsx_path


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("Uso: python fgts_extrato_to_excel.py <extrato.txt> [saida.xlsx]")
        raise SystemExit(1)

    txt = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) >= 3 else "Extrato_FGTS_Analitico_Processado.xlsx"
    generated = extrato_fgts_txt_para_excel(txt, out)
    print(f"Planilha gerada em: {generated}")
